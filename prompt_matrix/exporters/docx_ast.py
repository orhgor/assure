"""JDF AST to OpenXML (.docx) headless compiler."""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TRUST_BLUE = RGBColor(0x1A, 0x4B, 0x8C)
SECTION_COLOR = RGBColor(0x0D, 0x2B, 0x45)
HEADER_FILL = "1A4B8C"
CALLOUT_COLORS = {
    "adversarial_redhat": "FFF1F2",
    "warning": "FEF3C7",
    "insight": "EFF6FF",
}
_NUMERIC_RE = re.compile(r"^-?\d+(?:\.\d+)?%?$")


def _set_cell_shading(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _is_numeric(value: str) -> bool:
    return bool(_NUMERIC_RE.match((value or "").strip()))


def _add_paragraph(doc: Document, text: str) -> None:
    para = doc.add_paragraph(text or "")
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    para.paragraph_format.line_spacing = 1.15
    para.paragraph_format.space_after = Pt(8)
    for run in para.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)


def _add_callout(doc: Document, node: dict[str, Any]) -> None:
    variant = str(node.get("variant") or "warning")
    fill = CALLOUT_COLORS.get(variant, "FEF3C7")
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _set_cell_shading(cell, fill)
    title = str(node.get("title") or variant.replace("_", " ").upper())
    body = str(node.get("content") or "")
    p_title = cell.paragraphs[0]
    run_title = p_title.add_run(title)
    run_title.bold = True
    run_title.font.color.rgb = TRUST_BLUE
    run_title.font.size = Pt(10)
    p_body = cell.add_paragraph(body)
    p_body.paragraph_format.space_after = Pt(4)
    for run in p_body.runs:
        run.italic = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)


def _add_table(doc: Document, node: dict[str, Any]) -> None:
    caption = str(node.get("caption") or "").strip()
    if caption:
        cap = doc.add_paragraph(caption)
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(10)

    headers = [str(h) for h in (node.get("headers") or [])]
    rows = node.get("rows") or []
    if not headers and not rows:
        return

    col_count = max(len(headers), max((len(r) for r in rows), default=0))
    if col_count == 0:
        return

    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.style = "Table Grid"

    if headers:
        for idx, header in enumerate(headers):
            cell = table.rows[0].cells[idx]
            cell.text = header
            _set_cell_shading(cell, HEADER_FILL)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.size = Pt(10)

    start_row = 1 if headers else 0
    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            val = str(row[c_idx]) if c_idx < len(row) else ""
            cell = table.rows[start_row + r_idx].cells[c_idx]
            cell.text = val
            if _is_numeric(val):
                for paragraph in cell.paragraphs:
                    paragraph.alignment = 2  # RIGHT


def _add_truth_appendix(doc: Document, ledger: dict[str, Any]) -> None:
    doc.add_page_break()
    heading = doc.add_heading("Appendix: Truth Ledger Receipts", level=1)
    for run in heading.runs:
        run.font.name = "Georgia"
        run.font.color.rgb = SECTION_COLOR

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Canonical key", "Locked value", "Verification status"]
    for idx, label in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = label
        _set_cell_shading(cell, HEADER_FILL)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for key, value in sorted(ledger.items(), key=lambda item: str(item[0])):
        row = table.add_row().cells
        row[0].text = str(key)
        row[1].text = str(value)
        row[2].text = "LOCKED"


def export_jdf_to_docx(jdf_tree: dict[str, Any]) -> io.BytesIO:
    """Compile a JDF AST dict into a styled Word document buffer."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    meta = jdf_tree.get("meta") or {}
    title = str(meta.get("title") or meta.get("project_id") or jdf_tree.get("document_id") or "Assure Document")

    title_para = doc.add_paragraph()
    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = TRUST_BLUE
    title_para.paragraph_format.space_after = Pt(12)

    for section in jdf_tree.get("body") or []:
        if not isinstance(section, dict):
            continue
        sec_title = str(section.get("title") or "Section")
        sec_heading = doc.add_heading(sec_title, level=2)
        for run in sec_heading.runs:
            run.font.name = "Georgia"
            run.font.size = Pt(14)
            run.font.color.rgb = SECTION_COLOR

        for child in section.get("children") or []:
            if not isinstance(child, dict):
                continue
            ntype = child.get("type")
            if ntype == "paragraph":
                _add_paragraph(doc, str(child.get("content") or ""))
            elif ntype == "callout":
                _add_callout(doc, child)
            elif ntype == "table":
                _add_table(doc, child)

    ledger = jdf_tree.get("truth_ledger") or {}
    if ledger:
        _add_truth_appendix(doc, ledger)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
