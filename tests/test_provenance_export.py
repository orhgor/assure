"""Provenance parsing and DOCX citation export tests."""

from __future__ import annotations

import io
import tempfile
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

from prompt_matrix.exporters.docx_ast import export_jdf_to_docx
from prompt_matrix.models.jdf import (
    collect_unique_provenance,
    extract_citations_from_content,
    parse_document,
)


def test_extract_citations_from_content():
    content = (
        'Revenue reached $12M per <cite id="cite-1" source_type="internal_doc" '
        'source_name="cim.pdf" page_number="12">Revenue was $12M</cite> filing.'
    )
    cleaned, prov = extract_citations_from_content(content)
    assert "cite" not in cleaned
    assert "$12M" in cleaned
    assert len(prov) == 1
    assert prov[0]["source_id"] == "cite-1"
    assert prov[0]["source_name"] == "cim.pdf"
    assert prov[0]["extracted_quote"] == "Revenue was $12M"


def test_parse_document_enriches_citations():
    raw = {
        "document_id": "doc-test",
        "meta": {},
        "truth_ledger": {},
        "body": [
            {
                "type": "section",
                "id": "sec-1",
                "title": "Summary",
                "children": [
                    {
                        "type": "paragraph",
                        "id": "p-1",
                        "content": (
                            'Claim <cite id="c1" source_type="web_url" source_name="Example" '
                            'url_or_doi="https://example.com">quoted</cite>.'
                        ),
                        "entities_referenced": [],
                        "meta": {},
                        "annotations": {"redhat": [], "z3": []},
                    }
                ],
                "meta": {},
                "annotations": {"redhat": [], "z3": []},
            }
        ],
    }
    doc = parse_document(raw)
    para = doc.body[0].children[0]
    assert "cite" not in para.content
    assert len(para.provenance) == 1
    assert para.provenance[0].source_type == "web_url"
    assert para.provenance[0].url_or_doi == "https://example.com"


def test_migrate_legacy_provenance_object():
    raw = {
        "document_id": "doc-legacy",
        "meta": {},
        "truth_ledger": {},
        "body": [
            {
                "type": "section",
                "id": "sec-1",
                "title": "T",
                "children": [
                    {
                        "type": "paragraph",
                        "id": "p-1",
                        "content": "Text",
                        "provenance": [
                            {
                                "source_file": "old.pdf",
                                "page_or_timestamp": "3",
                                "exact_quote": "hello",
                            }
                        ],
                        "entities_referenced": [],
                        "meta": {},
                        "annotations": {"redhat": [], "z3": []},
                    }
                ],
                "meta": {},
                "annotations": {"redhat": [], "z3": []},
            }
        ],
    }
    doc = parse_document(raw)
    prov = doc.body[0].children[0].provenance[0]
    assert prov.source_name == "old.pdf"
    assert prov.page_number == "3"
    assert prov.extracted_quote == "hello"
    assert prov.source_type == "internal_doc"


def test_collect_unique_provenance_dedupes():
    tree = {
        "body": [
            {
                "children": [
                    {
                        "type": "paragraph",
                        "id": "p1",
                        "provenance": [{"source_id": "a", "source_name": "Doc A"}],
                    },
                    {
                        "type": "paragraph",
                        "id": "p2",
                        "provenance": [{"source_id": "a", "source_name": "Doc A duplicate"}],
                    },
                ]
            }
        ]
    }
    refs = collect_unique_provenance(tree)
    assert len(refs) == 1


def test_export_docx_includes_references_by_default():
    tree = {
        "document_id": "doc-export",
        "meta": {"title": "Export Test"},
        "truth_ledger": {},
        "body": [
            {
                "type": "section",
                "id": "sec-1",
                "title": "Body",
                "children": [
                    {
                        "type": "paragraph",
                        "id": "p-1",
                        "content": "Revenue grew.",
                        "provenance": [
                            {
                                "source_type": "internal_doc",
                                "source_name": "Board deck",
                                "source_id": "ref-1",
                                "page_number": "4",
                                "extracted_quote": "Revenue grew 20%",
                            }
                        ],
                    }
                ],
            }
        ],
    }
    buffer = export_jdf_to_docx(tree, include_citations=True)
    doc = Document(io.BytesIO(buffer.getvalue()))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "References" in text
    assert "Board deck" in text


def test_export_docx_skips_references_when_disabled():
    tree = {
        "document_id": "doc-export",
        "meta": {"title": "Export Test"},
        "truth_ledger": {},
        "body": [
            {
                "type": "section",
                "id": "sec-1",
                "title": "Body",
                "children": [
                    {
                        "type": "paragraph",
                        "id": "p-1",
                        "content": "Revenue grew.",
                        "provenance": [
                            {
                                "source_type": "internal_doc",
                                "source_name": "Board deck",
                                "source_id": "ref-1",
                            }
                        ],
                    }
                ],
            }
        ],
    }
    buffer = export_jdf_to_docx(tree, include_citations=False)
    doc = Document(io.BytesIO(buffer.getvalue()))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "References" not in text


@pytest.fixture
def settings_client():
    tmp = tempfile.TemporaryDirectory()
    db_path = Path(tmp.name) / "history.sqlite"

    def _getter():
        import sqlite3

        conn = sqlite3.connect(str(db_path), check_same_thread=False)
        conn.row_factory = sqlite3.Row
        return conn

    with (
        patch("prompt_matrix.history.DB_PATH", db_path),
        patch("prompt_matrix.db.connection.get_db", side_effect=_getter),
        patch("prompt_matrix.db.jdf_repository.get_db", side_effect=_getter),
        patch("prompt_matrix.db.settings_repository.get_db", side_effect=_getter),
    ):
        from prompt_matrix.db.connection import init_db
        from prompt_matrix.web import create_app

        init_db(_getter())
        client = create_app(require_auth=False).test_client()
        yield client
        tmp.cleanup()


def test_project_settings_roundtrip(settings_client):
    res = settings_client.put(
        "/api/projects/default/settings",
        json={"show_citations": False},
    )
    assert res.status_code == 200
    body = res.get_json()
    assert body["settings"]["show_citations"] is False

    got = settings_client.get("/api/projects/default/settings")
    assert got.status_code == 200
    assert got.get_json()["settings"]["show_citations"] is False
